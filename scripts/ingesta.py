#!/usr/bin/env python3
"""
Convierte documentos a texto plano para que Benjamin los pueda leer.

    python scripts/ingesta.py                      # convierte lo que haya suelto en knowledge/
    python scripts/ingesta.py ~/Documentos/*.pdf   # convierte archivos concretos

Deja un .md por documento dentro de knowledge/. El agente solo lee .md y .txt, a
proposito: el contexto queda en texto revisable y versionado en git, no en binarios.

Formatos: .pdf, .docx, .csv, .json, .yaml, .yml, .html, .htm, .md, .txt
Para PDF y Word hacen falta dos librerias:

    pip install pypdf python-docx
"""

import csv
import io
import json
import re
import sys
import unicodedata
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
KNOWLEDGE = RAIZ / "knowledge"
YA_LEGIBLES = (".md", ".txt")


def slug(texto: str) -> str:
    """Nombre de archivo en minusculas, sin acentos ni espacios."""
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    texto = re.sub(r"[^\w\s-]", "", texto).strip().lower()
    return re.sub(r"[-\s]+", "-", texto) or "documento"


def de_pdf(ruta: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("falta pypdf. Instalalo con: pip install pypdf")

    lector = PdfReader(str(ruta))
    paginas = []
    for i, pagina in enumerate(lector.pages, 1):
        texto = (pagina.extract_text() or "").strip()
        if texto:
            paginas.append(f"### Pagina {i}\n\n{texto}")

    if not paginas:
        raise RuntimeError(
            "el PDF no tiene texto extraible. Si es un escaneo, hay que pasarle OCR primero"
        )
    return "\n\n".join(paginas)


def de_docx(ruta: Path) -> str:
    try:
        import docx
    except ImportError:
        raise RuntimeError("falta python-docx. Instalalo con: pip install python-docx")

    documento = docx.Document(str(ruta))
    partes = []

    for parrafo in documento.paragraphs:
        texto = parrafo.text.strip()
        if not texto:
            continue
        estilo = (parrafo.style.name or "").lower()
        if estilo.startswith("heading"):
            nivel = "".join(c for c in estilo if c.isdigit()) or "2"
            partes.append(f"{'#' * min(int(nivel) + 1, 6)} {texto}")
        else:
            partes.append(texto)

    for tabla in documento.tables:
        filas = [[celda.text.strip() for celda in fila.cells] for fila in tabla.rows]
        if filas:
            partes.append(tabla_markdown(filas))

    if not partes:
        raise RuntimeError("el documento no tiene texto")
    return "\n\n".join(partes)


def tabla_markdown(filas: list[list[str]]) -> str:
    """Pasa una lista de filas a tabla markdown. Claude lee tablas bien."""
    if not filas:
        return ""
    ancho = max(len(f) for f in filas)
    filas = [f + [""] * (ancho - len(f)) for f in filas]
    cabecera = filas[0]
    lineas = [
        "| " + " | ".join(cabecera) + " |",
        "|" + "|".join("---" for _ in cabecera) + "|",
    ]
    for fila in filas[1:]:
        lineas.append("| " + " | ".join(c.replace("|", "\\|") for c in fila) + " |")
    return "\n".join(lineas)


def de_csv(ruta: Path) -> str:
    texto = ruta.read_text(encoding="utf-8", errors="replace")
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096])
    except csv.Error:
        dialecto = csv.excel
    filas = [f for f in csv.reader(io.StringIO(texto), dialecto) if any(c.strip() for c in f)]
    if not filas:
        raise RuntimeError("el CSV esta vacio")
    return tabla_markdown(filas)


def de_json(ruta: Path) -> str:
    datos = json.loads(ruta.read_text(encoding="utf-8"))
    return "```json\n" + json.dumps(datos, indent=2, ensure_ascii=False) + "\n```"


def de_yaml(ruta: Path) -> str:
    import yaml

    datos = yaml.safe_load(ruta.read_text(encoding="utf-8"))
    return "```yaml\n" + yaml.safe_dump(datos, allow_unicode=True, sort_keys=False) + "```"


def de_html(ruta: Path) -> str:
    bruto = ruta.read_text(encoding="utf-8", errors="replace")
    bruto = re.sub(r"(?is)<(script|style).*?</\1>", " ", bruto)
    texto = re.sub(r"(?s)<[^>]+>", " ", bruto)
    texto = re.sub(r"&nbsp;?", " ", texto)
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n\s*\n\s*\n+", "\n\n", texto).strip()
    if not texto:
        raise RuntimeError("el HTML no tiene texto")
    return texto


CONVERSORES = {
    ".pdf": de_pdf,
    ".docx": de_docx,
    ".csv": de_csv,
    ".json": de_json,
    ".yaml": de_yaml,
    ".yml": de_yaml,
    ".html": de_html,
    ".htm": de_html,
}


def siguiente_numero() -> int:
    """El prefijo numerico que toca, para que el orden de lectura sea estable."""
    usados = []
    for ruta in KNOWLEDGE.glob("*.md"):
        m = re.match(r"^(\d+)-", ruta.name)
        if m:
            usados.append(int(m.group(1)))
    return max(usados, default=0) + 1


def convertir(ruta: Path, numero: int) -> Path | None:
    extension = ruta.suffix.lower()

    if extension in YA_LEGIBLES and ruta.parent == KNOWLEDGE:
        print(f"  {ruta.name}: ya es legible, se deja como esta")
        return None

    if extension in YA_LEGIBLES:
        texto = ruta.read_text(encoding="utf-8", errors="replace")
    else:
        conversor = CONVERSORES.get(extension)
        if conversor is None:
            print(f"  {ruta.name}: formato no soportado ({extension}), se salta")
            return None
        try:
            texto = conversor(ruta)
        except Exception as e:  # noqa: BLE001
            print(f"  {ruta.name}: ERROR, {e}")
            return None

    destino = KNOWLEDGE / f"{numero:02d}-{slug(ruta.stem)}.md"
    if destino.exists():
        print(f"  {ruta.name}: ya existe {destino.name}, se salta. Borralo si querés rehacerlo")
        return None

    destino.write_text(
        f"# {ruta.stem}\n\n_Convertido desde {ruta.name}_\n\n{texto.strip()}\n",
        encoding="utf-8",
    )
    print(f"  {ruta.name}  ->  {destino.name}  ({len(texto):,} caracteres)")
    return destino


def main(argumentos: list[str]) -> int:
    KNOWLEDGE.mkdir(exist_ok=True)

    if argumentos:
        fuentes = [Path(a).expanduser() for a in argumentos]
        faltantes = [f for f in fuentes if not f.is_file()]
        if faltantes:
            for f in faltantes:
                print(f"No existe: {f}")
            return 1
    else:
        fuentes = sorted(
            r for r in KNOWLEDGE.iterdir()
            if r.is_file()
            and not r.name.startswith(".")
            and r.suffix.lower() not in YA_LEGIBLES
        )
        if not fuentes:
            print(f"No hay nada que convertir en {KNOWLEDGE}/.")
            print("Pasa archivos como argumento: python scripts/ingesta.py documento.pdf")
            return 0

    print(f"Convirtiendo {len(fuentes)} archivo(s) a texto en knowledge/:\n")
    numero = siguiente_numero()
    convertidos = 0
    for fuente in fuentes:
        if convertir(fuente, numero):
            numero += 1
            convertidos += 1

    print(f"\n{convertidos} archivo(s) convertidos.")
    if convertidos:
        print("Revisa el texto antes de confiar en el: la extraccion nunca es perfecta.")
        print("Despues corre: python scripts/contexto.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
